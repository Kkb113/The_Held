"""Read-only audit of the delivered v3 package; reports go beside this script."""
from pathlib import Path
import csv, hashlib, importlib.util, json, re, zipfile
from collections import Counter
from unittest.mock import patch
from xml.etree import ElementTree as ET
from pypdf import PdfReader
import openpyxl

OUT = Path(__file__).resolve().parent
ROOT = OUT.parents[1]
PACKAGE = ROOT / 'blueprint/THE_HELD_First_Step_v3'
def digest(p): return hashlib.sha256(p.read_bytes()).hexdigest()
manifest = json.loads((PACKAGE/'MANIFEST.json').read_text(encoding='utf-8'))
integrity = []
for row in manifest['files']:
    p = PACKAGE / row['path']
    integrity.append({'path': row['path'], 'pass': p.stat().st_size == row['bytes'] and digest(p) == row['sha256']})
assert all(r['pass'] for r in integrity), integrity
data = {p.stem: json.loads(p.read_text(encoding='utf-8')) for p in (PACKAGE/'registers').glob('*.json')}
mirrors = {}
for name, rows in data.items():
    with (PACKAGE/'registers'/f'{name}.csv').open(encoding='utf-8-sig',newline='') as f:
        exported = list(csv.DictReader(f))
    assert [r['id'] for r in rows] == [r['id'] for r in exported], name
    mirrors[name] = {'rows': len(rows), 'ordered_ids_match': True}

# The supplied validator writes one report; redirect that write to the audit folder
# so the delivered package and original manifest remain byte-for-byte unchanged.
spec = importlib.util.spec_from_file_location('held_v3_validator', PACKAGE/'validate_package.py')
module = importlib.util.module_from_spec(spec)
spec.loader.exec_module(module)
original_write = Path.write_text
def report_write(p, text, *args, **kwargs):
    assert p == PACKAGE/'validation_report.json', p
    return original_write(OUT/'planning_validation.json', text, *args, **kwargs)
with patch.object(Path, 'write_text', report_write):
    module.main()

master = (PACKAGE/'THE_HELD_First_Step_Master_v3.md').read_text(encoding='utf-8')
with zipfile.ZipFile(PACKAGE/'THE_HELD_First_Step_Master_v3.docx') as z:
    xml = ET.fromstring(z.read('word/document.xml'))
    ns = {'w':'http://schemas.openxmlformats.org/wordprocessingml/2006/main'}
    docx_text = '\n'.join(''.join(p.itertext()) for p in xml.findall('.//w:t', ns))
pdf = PdfReader(PACKAGE/'THE_HELD_First_Step_Master_v3.pdf')
pdf_text = '\n'.join(p.extract_text() or '' for p in pdf.pages)
sequence_ids = [s['id'] for s in data['sequences']]
document_presence = {fmt: all(s in content for s in sequence_ids) for fmt, content in [('markdown', master),('docx', docx_text),('pdf',pdf_text)]}
assert all(document_presence.values())
assert len(pdf.pages) == 120
bookpath = PACKAGE/'THE_HELD_First_Step_Production_Registers_v3.xlsx'
book = openpyxl.load_workbook(bookpath, read_only=True, data_only=False)
cached = openpyxl.load_workbook(bookpath, read_only=True, data_only=True)
sheet_names = {'chapters':'Chapters','sequences':'Sequences','puzzles':'Puzzles','assets':'Assets','systems':'Systems','animation':'Animation','audio':'Audio','milestones':'Milestones','tests':'Tests','risks':'Risks','sprint01':'Sprint01','cohort':'Cohort','narrative_facts':'NarrativeFacts','sources':'Sources'}
workbook_ids = {}
for name, sheet in sheet_names.items():
    values = {str(c.value) for row in book[sheet] for c in row if c.value is not None}
    missing = [r['id'] for r in data[name] if r['id'] not in values]
    assert not missing, (sheet, missing)
    workbook_ids[sheet] = {'required_ids': len(data[name]),'missing':missing}
errors = [(s.title,c.coordinate,c.value) for s in cached for row in s for c in row if c.data_type == 'e']
formulas = [(s.title,c.coordinate,c.value) for s in book for row in s for c in row if c.data_type == 'f']
assert not errors
report = {'scope':'Document and planning-data audit only. No engine tests or new playtests executed.',
          'package_integrity':integrity, 'mirrors':mirrors,
          'master_lines':len(master.splitlines()), 'master_words':len(master.split()),
          'pdf_pages':len(pdf.pages),'all_sequence_ids_present_in_documents':document_presence,
          'workbook_sheets':book.sheetnames,'workbook_ids':workbook_ids,
          'workbook_formula_count':len(formulas),'workbook_cached_errors':errors,
          'workbook_limit':'Read-only IDs and cached-error checks; no native Excel recalculation or cell-by-cell prose equality claim.',
          'animation_count':len(data['animation']), 'audio_variants':sum(a['initial_variants'] for a in data['audio']),
          'game_test_statuses':dict(Counter(t['status'] for t in data['tests'])),
          'document_limit':'Master Markdown read in full; PDF/DOCX extraction verifies sequence coverage, not exact typography/prose equivalence.',
          'source_zip_sha256':digest(Path(r'C:\Users\karth\Downloads\THE_HELD_First_Step_Production_Package_v3.zip'))}
(OUT/'package_audit.json').write_text(json.dumps(report,indent=2),encoding='utf-8')
assert all(digest(PACKAGE/r['path']) == r['sha256'] for r in manifest['files'])
print(json.dumps({k:v for k,v in report.items() if k not in ['package_integrity','workbook_ids','mirrors']},indent=2))
