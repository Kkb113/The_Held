from pathlib import Path
import json, csv, hashlib, re, zipfile
from xml.etree import ElementTree as ET
from pypdf import PdfReader
import openpyxl

root = Path(__file__).resolve().parents[1]
out = root / 'review' / 'audit'
out.mkdir(exist_ok=True)
manifest = json.loads((root/'PACKAGE_MANIFEST.json').read_text(encoding='utf-8'))
audit = {'hashes': [], 'registers': {}, 'workbook': {}, 'documents': {}}
for item in manifest['files']:
    data = (root/item['path']).read_bytes()
    audit['hashes'].append({'file': item['path'], 'matches': len(data)==item['bytes'] and hashlib.sha256(data).hexdigest()==item['sha256']})
registers = {}
for p in sorted((root/'registers').glob('*.json')):
    records = json.loads(p.read_text(encoding='utf-8-sig'))
    registers[p.stem] = records
    with p.with_suffix('.csv').open(encoding='utf-8-sig', newline='') as f:
        rows = list(csv.DictReader(f))
    mismatches=[]
    for i,(a,b) in enumerate(zip(records,rows)):
        for key,value in a.items():
            v = '; '.join(map(str,value)) if isinstance(value,list) else str(value)
            if b.get(key)!=v: mismatches.append({'row':i+1,'key':key,'json':value,'csv':b.get(key)})
    audit['registers'][p.stem]={'json_rows':len(records),'csv_rows':len(rows),'mismatches':mismatches}
    (out/(p.stem+'.txt')).write_text('\n'.join(json.dumps(r,ensure_ascii=False) for r in records),encoding='utf-8')
wb = openpyxl.load_workbook(root/'THE_HELD_Production_Registers_v1.xlsx',data_only=False)
cached = openpyxl.load_workbook(root/'THE_HELD_Production_Registers_v1.xlsx',data_only=True)
for ws in wb:
    values=[]; formulas=[]; errors=[]
    for row in ws:
        cells=[]
        for c in row:
            if c.value is None: continue
            v=cached[ws.title][c.coordinate].value
            cells.append(f'{c.coordinate}: {c.value}'+(f' [cached={v}]' if c.data_type=='f' else ''))
            if c.data_type=='f': formulas.append({'cell':c.coordinate,'formula':c.value,'cached':v})
            if c.data_type=='e' or cached[ws.title][c.coordinate].data_type=='e': errors.append(c.coordinate)
        if cells: values.append(' | '.join(cells))
    (out/('workbook_'+ws.title.replace(' ','_')+'.txt')).write_text('\n'.join(values),encoding='utf-8')
    audit['workbook'][ws.title]={'rows':ws.max_row,'cols':ws.max_column,'formulas':formulas,'errors':errors,'images':len(ws._images),'charts':len(ws._charts)}
with zipfile.ZipFile(root/'THE_HELD_Master_Blueprint_v1.docx') as z:
    tree=ET.fromstring(z.read('word/document.xml'))
    ns={'w':'http://schemas.openxmlformats.org/wordprocessingml/2006/main'}
    # Text nodes preserve paragraph and table order.
    paragraphs=[''.join(t.text or '' for t in p.findall('.//w:t',ns)) for p in tree.findall('.//w:p',ns)]
    (out/'docx_text.txt').write_text('\n'.join(paragraphs),encoding='utf-8')
    audit['documents']['docx']={'paragraphs':len(paragraphs),'embedded_media':[n for n in z.namelist() if n.startswith('word/media/')],'characters':sum(map(len,paragraphs))}
pdf=PdfReader(root/'THE_HELD_Master_Blueprint_v1.pdf')
pdftexts=[p.extract_text() or '' for p in pdf.pages]
(out/'pdf_text.txt').write_text('\n\n'.join(f'PAGE {i+1}\n{t}' for i,t in enumerate(pdftexts)),encoding='utf-8')
audit['documents']['pdf']={'pages':len(pdf.pages),'image_counts':[len(p.images) for p in pdf.pages],'characters':sum(map(len,pdftexts))}
md=(root/'THE_HELD_Master_Blueprint_v1.md').read_text(encoding='utf-8')
def norm(s): return re.sub(r'[^\w]','',s).lower()
mdn=norm(md)
docx_extra=[p for p in paragraphs if len(norm(p))>40 and norm(p) not in mdn]
audit['documents']['docx_paragraphs_not_in_md']=docx_extra
seq=registers['sequences']
audit['totals']={'sequence_minutes':sum(r['minutes'] for r in seq),'chapter_minutes':sum(r['minutes'] for r in registers['chapters']),'animation':sum(r['allocation'] for r in registers['animation']),'audio':sum(r['allocation'] for r in registers['audio']),'environment':sum(r['allocation'] for r in registers['environment_kits'])}
audit['sequence_schema']=list(seq[0])
audit['asset_schema']=list(registers['assets'][0])
seq_mismatch=[]
for r in seq:
    for k in ('title','space','entry','action','threat','exit','checkpoint','test'):
        if norm(r[k]) not in mdn: seq_mismatch.append([r['id'],k])
audit['sequence_fields_not_in_md']=seq_mismatch
pdfbody=' '.join(line for t in pdftexts for line in t.splitlines() if not line.startswith('THE HELD    /') and not line.startswith('VERSION 1.0'))
pdfn=norm(pdfbody)
audit['docx_body_paragraphs_not_in_pdf']=[p for p in paragraphs if len(norm(p))>40 and norm(p) not in pdfn]
maps={
 'Chapters':('chapters',['id','title','minutes',None,None,'goal',None]),
 'Sequences':('sequences',['id','chapter','title','minutes','campaign_start_min','campaign_end_min','kind','space','assets',None]),
 'Assets':('assets',['id','name','type','owner','first_use','deliverable','gate','status']),
 'Systems':('systems',['id','name','type','owner','first_use','deliverable','gate','status']),
 'Animation':('animation',['group','allocation','unit','coverage','status']),
 'Audio':('audio',['group','allocation','unit','coverage','status']),
 'Environment Kits':('environment_kits',['group','allocation','unit','coverage','status']),
 'Milestones':('milestones',['id','name','start_week','end_week',None,'dependencies','owner','playable','gate','status']),
 'Tests':('tests',['id','area','setup','action','expected','gate','method','status']),
 'Risks':('risks',['id','risk','likelihood','impact','owner','mitigation','review_gate','fallback','status']),
 'First Sprint':('first_sprint',['id','day','owner','task','output','dependency','status']),
 'Sources':('sources',['id','publisher','title','url','use','accessed']),
}
audit['workbook_register_comparison']={}
for name,(reg,fields) in maps.items():
    discrepancies=[]
    for i,r in enumerate(registers[reg],start=5):
        for j,key in enumerate(fields,start=1):
            if key is None or key not in r: continue
            rv=r[key]
            if isinstance(rv,list): rv='; '.join(map(str,rv))
            cv=cached[name].cell(i,j).value
            if str(rv)!=str(cv): discrepancies.append({'cell':cached[name].cell(i,j).coordinate,'key':key,'register':rv,'workbook':cv})
    audit['workbook_register_comparison'][name]=discrepancies
(out/'package_audit.json').write_text(json.dumps(audit,ensure_ascii=False,indent=2),encoding='utf-8')
print(json.dumps({'hashes_pass':sum(x['matches'] for x in audit['hashes']),'register_rows':sum(len(r) for r in registers.values()),'json_csv_mismatch_count':sum(len(r['mismatches']) for r in audit['registers'].values()),'totals':audit['totals'],'sequence_fields_not_in_md':seq_mismatch,'docx_pdf_differences':audit['docx_body_paragraphs_not_in_pdf'],'workbook_differences':audit['workbook_register_comparison']},ensure_ascii=False,indent=2))
